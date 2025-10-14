import openpyxl
import os

OUTPUT_DIR = 'output_timetables'
HTML_DIR = '.'

BRANCHES = ['CSE', 'DSAI', 'ECE']
SEMESTERS = [1, 3, 5, 7]

def excel_to_html_table(ws):
    html = '<table>\n<thead>\n<tr>'
    for cell in ws[1]:
        html += f'<th>{cell.value}</th>'
    html += '</tr>\n</thead>\n<tbody>'
    for row in ws.iter_rows(min_row=2, values_only=True):
        html += '<tr>' + ''.join(f'<td>{cell if cell is not None else ""}</td>' for cell in row) + '</tr>'
    html += '</tbody>\n</table>'
    return html

def generate_html_files():
    links = []
    for semester in SEMESTERS:
        for branch in BRANCHES:
            filename = f"sem{semester}_{branch}_timetable.xlsx"
            sem_name = f"Semester {semester} - {branch}"
            path = os.path.join(OUTPUT_DIR, filename)
            if not os.path.exists(path):
                continue
            wb = openpyxl.load_workbook(path)
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                html_table = excel_to_html_table(ws)
            # ------------------ Modified Timetable HTML ------------------
            html_content = f'''<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{sem_name} - {sheet}</title>
<link href="https://fonts.googleapis.com/css?family=Poppins:400,600&display=swap" rel="stylesheet">
<style>
body {{
  font-family: 'Poppins', sans-serif;
  background: #f4f7fb;
  margin: 0;
  padding: 0;
}}
.container {{
  background: #fff;
  margin: 40px auto;
  padding: 32px;
  border-radius: 18px;
  box-shadow: 0 8px 32px rgba(0,0,0,0.12);
  max-width: 950px;
}}
h1 {{
  color: #4a2c82;
  text-align: center;
  margin-bottom: 24px;
  font-weight: 600;
}}
table {{
  width: 100%;
  border-collapse: collapse;
  margin-top: 16px;
  font-size: 15px;
}}
th, td {{
  border: 1px solid #d0d7de;
  padding: 12px 10px;
  text-align: center;
}}
th {{
  background: linear-gradient(135deg,#4a2c82,#3bb273);
  color: #fff;
}}
tr:nth-child(even) {{ background: #f9f9ff; }}
tr:hover {{ background: #eef3ff; }}
a {{
  display: inline-block;
  margin-bottom: 16px;
  color: #4a2c82;
  text-decoration: none;
  font-weight: bold;
}}
a:hover {{ text-decoration: underline; color: #3bb273; }}
</style>
</head>
<body>
<div class="container">
<a href="index.html">&#8592; Back to Timetable Selection</a>
<h1>{sem_name} - {sheet}</h1>
{html_table}
</div>
</body>
</html>'''
            out_name = f"{filename.replace('.xlsx', '')}_{sheet.replace(' ', '_')}.html"
            links.append((sem_name, sheet, out_name))
            with open(os.path.join(HTML_DIR, out_name), 'w', encoding='utf-8') as f:
                f.write(html_content)

    # ------------------ Modified Index Page ------------------
    sem_data = {}
    for sem_name, sheet, out_name in links:
        sem_data.setdefault(sem_name, []).append((sheet, out_name))

    index_html = '''<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Timetable Selector</title>
<link href="https://fonts.googleapis.com/css?family=Poppins:400,600&display=swap" rel="stylesheet">
<style>
body {
  font-family: 'Poppins', sans-serif;
  background: linear-gradient(120deg, #4a2c82, #3bb273);
  margin: 0;
  padding: 0;
  text-align: center;
}
.container {
  background: #fff;
  margin: 40px auto;
  padding: 32px;
  border-radius: 18px;
  box-shadow: 0 8px 32px rgba(0,0,0,0.12);
  max-width: 950px;
}
h1 { 
  color: #333; 
  margin-bottom: 24px; 
  font-weight: 600;
}
select {
  padding: 12px 15px;
  margin: 12px;
  border-radius: 10px;
  border: 2px solid #4a2c82;
  font-size: 15px;
  background: #f9f9f9;
  cursor: pointer;
  transition: 0.3s;
}
select:hover {
  border-color: #3bb273;
}
iframe {
  width: 100%;
  height: 650px;
  border: none;
  border-radius: 14px;
  margin-top: 25px;
  background: #fdfdfd;
  box-shadow: inset 0 0 10px rgba(0,0,0,0.1);
}
</style>
</head>
<body>
<div class="container">
<h1>📘 Timetable Selector</h1>
<label for="semester"><b>Semester:</b></label>
<select id="semester" onchange="updateSections()">
  <option value="">-- Choose Semester --</option>
</select>
<label for="section"><b>Section:</b></label>
<select id="section" onchange="showTimetable()">
  <option value="">-- Choose Section --</option>
</select>
<iframe id="timetableFrame" src=""></iframe>
</div>
<script>
const semData = {
'''
    for sem_name, sheets in sem_data.items():
        index_html += f'  "{sem_name}": {{\n'
        for sheet, out_name in sheets:
            index_html += f'    "{sheet}": "{out_name}",\n'
        index_html += '  },\n'
    index_html += '''};

const semesterSelect = document.getElementById('semester');
const sectionSelect = document.getElementById('section');
const frame = document.getElementById('timetableFrame');

// Populate semester dropdown
Object.keys(semData).forEach(sem => {
  const opt = document.createElement('option');
  opt.value = sem;
  opt.textContent = sem;
  semesterSelect.appendChild(opt);
});

function updateSections() {
  const selectedSem = semesterSelect.value;
  sectionSelect.innerHTML = '<option value="">-- Choose Section --</option>';
  if (selectedSem && semData[selectedSem]) {
    Object.keys(semData[selectedSem]).forEach(section => {
      const opt = document.createElement('option');
      opt.value = section;
      opt.textContent = section;
      sectionSelect.appendChild(opt);
    });
  }
  frame.src = "";
}

function showTimetable() {
  const sem = semesterSelect.value;
  const sec = sectionSelect.value;
  if (sem && sec && semData[sem][sec]) {
    frame.src = semData[sem][sec];
  }
}
</script>
</body>
</html>'''
    with open(os.path.join(HTML_DIR, 'index.html'), 'w', encoding='utf-8') as f:
        f.write(index_html)
